#!/usr/bin/env python3
"""Build the deal workbook from deals.json.

Tabs: Summary, Top 10, Pick, Rubric, Deals, Market Sizing, and, when the data exists, Second Scorer,
Backtest, Fund Comps, Form D Leads and Sources. Scores, ranks, TAMs and the max price are live formulas:
change a weight on Rubric or an assumption on Pick and everything updates. Blue text = inputs,
black = formulas, green = links to other tabs, yellow fill = assumptions.

Usage: python3 build_workbook.py deals.json [--out <Fund>_Deals.xlsx]
"""
import argparse
import math
from copy import copy
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from _common import load, weights, KEYS, LABELS, DEFAULT_WEIGHTS, slug

F = "Arial"
NAVY = "1F3A5F"
f_base = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True, color=NAVY)
f_sub = Font(name=F, size=10, italic=True, color="555555")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_in = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
f_link_b = Font(name=F, size=10, color="008000", bold=True)
f_url = Font(name=F, size=10, color="0563C1", underline="single")
fill_head = PatternFill("solid", fgColor=NAVY)
fill_key = PatternFill("solid", fgColor="FFFF00")
fill_band = PatternFill("solid", fgColor="EEF2F7")
thin = Side(style="thin", color="C9D1DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
wrap_c = Alignment(wrap_text=True, vertical="top", horizontal="center")
right = Alignment(horizontal="right", vertical="top")
USD_M = '$#,##0.0"M";($#,##0.0"M");-'
USD = '$#,##0;($#,##0);-'
PCT = '0.0%;(0.0%);-'
SC = '0.00'
MULT = '0.0"x"'

RUBRIC_TEXT = {
    "D": ("Founder", "No prior exposure to the problem", "A few years adjacent to the problem",
          "10+ years inside the problem, or ran the function at scale"),
    "X": ("Founder", "First job or first company", "Early employee or exec at a venture-backed company",
          "Founded and sold a company, or senior exec through a large exit"),
    "T": ("Founder", "No technical co-founder; product outsourced", "Competent product and engineering lead",
          "PhD, research faculty or principal engineer; AI in production"),
    "C": ("Founder", "Solo founder with gaps in key roles", "Two founders with one clear gap",
          "Complementary CEO, CTO and go-to-market, all full-time"),
    "M": ("Market", "Under $100M bottom-up", "$300M-$1B bottom-up", "Over $1B bottom-up with clear adjacencies"),
    "TR": ("Traction", "Pre-product or no customers named", "Paying pilots or named customers",
           "Marquee customers plus growth or retention data"),
    "F": ("Fund fit", "Outside geography or past stage", "In region but awkward model, stage or price",
          "In region and sectors; fund can be 1st or 2nd institutional money at its check size"),
    "TM": ("Timing", "Just raised a large round", "Next raise likely in 12-24 months", "Raising now and needs a lead"),
}


def put(ws, ref, val, font=f_base, fmt=None, align=wrap, fill=None):
    c = ws[ref]
    c.value = val
    c.font = font
    c.alignment = align
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def header(ws, row, labels, widths):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill, c.alignment, c.border = f_head, fill_head, wrap_c, box
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def boxrow(ws, r, ncols):
    for col in range(1, ncols + 1):
        ws.cell(row=r, column=col).border = box


def section(ws, r, label, cols="BCD"):
    put(ws, f"{cols[0]}{r}", label, f_head, fill=fill_head)
    for c in cols[1:]:
        ws[f"{c}{r}"].fill = fill_head


def wrap_and_fit(ws):
    merged = {}
    for rng in ws.merged_cells.ranges:
        tl = ws.cell(rng.min_row, rng.min_col).coordinate
        merged[tl] = sum((ws.column_dimensions[get_column_letter(c)].width or 8.43) for c in range(rng.min_col, rng.max_col + 1))
    maxc = max(min(ws.max_column, 14), 7)
    for r in (2, 3):
        for c0 in (1, 2):
            c = ws.cell(r, c0)
            if isinstance(c.value, str) and len(c.value) > 40 and c.coordinate not in merged:
                if all(ws.cell(r, k).value is None for k in range(c0 + 1, maxc + 1)):
                    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=maxc)
                    merged[c.coordinate] = sum((ws.column_dimensions[get_column_letter(k)].width or 8.43) for k in range(c0, maxc + 1))
    heights = {}
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                al = copy(c.alignment)
                al.wrap_text = True
                al.vertical = al.vertical or "top"
                c.alignment = al
                continue
            if not isinstance(v, str):
                continue
            nxt = ws.cell(c.row, c.column + 1).value if c.column < ws.max_column else None
            in_table = c.border is not None and c.border.left is not None and c.border.left.style is not None
            if c.coordinate not in merged and nxt is None and not in_table:
                al = copy(c.alignment)
                al.wrap_text = False
                c.alignment = al
                continue
            al = copy(c.alignment)
            al.wrap_text = True
            al.vertical = al.vertical or "top"
            c.alignment = al
            width = merged.get(c.coordinate, ws.column_dimensions[c.column_letter].width or 8.43)
            per_line = max(1.0, (width - 1) * 1.15)
            size = c.font.sz or 10
            lines = sum(max(1, math.ceil(len(p) * (size / 10) / per_line)) for p in v.split("\n"))
            heights[c.row] = max(heights.get(c.row, 15), lines * size * 1.32 + 4)
    for r, h in heights.items():
        ws.row_dimensions[r].height = min(409, round(h, 1))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deals")
    ap.add_argument("--out")
    a = ap.parse_args()
    data = load(a.deals)
    fund = data.get("fund") or {}
    fname = fund.get("name", "Fund")
    out = a.out or f"{slug(fname)}_Deals.xlsx"
    n = int(data.get("deal_count", 40))
    W = weights(data)
    comps_all = data.get("companies", [])
    deals = sorted(comps_all, key=lambda c: c.get("rank", 9999))[:n]
    top10 = deals[:10]
    pick = data.get("pick") or {}
    wb = Workbook()

    # ---------- Rubric ----------
    ru = wb.active
    ru.title = "Rubric"
    put(ru, "A1", "Scoring rubric and weights", f_title, align=Alignment())
    note = "Blue weights are inputs: change one and every score and rank updates."
    if data.get("original_weights"):
        note = "Weights tuned by backtesting the fund's past first checks (Backtest tab). " + note
    put(ru, "A2", note, f_sub, align=Alignment())
    header(ru, 4, ["Factor", "Sub-factor", "Weight", "Score of 1", "Score of 3", "Score of 5"], [12, 22, 9, 34, 34, 42])
    for i, k in enumerate(KEYS):
        r = 5 + i
        grp, s1, s3, s5 = RUBRIC_TEXT[k]
        put(ru, f"A{r}", grp, f_bold)
        put(ru, f"B{r}", LABELS[k])
        put(ru, f"C{r}", W[k], f_in, PCT, wrap_c, fill_key)
        put(ru, f"D{r}", s1)
        put(ru, f"E{r}", s3)
        put(ru, f"F{r}", s5)
        boxrow(ru, r, 6)
    put(ru, "B13", "Total (must equal 100%)", f_bold)
    put(ru, "C13", "=SUM(C5:C12)", f_bold, PCT, wrap_c)
    put(ru, "B14", "Founder factors", f_bold)
    put(ru, "C14", "=SUM(C5:C8)", f_bold, PCT, wrap_c)
    put(ru, "B15", "Check", f_bold)
    put(ru, "C15", '=IF(ABS(C13-1)<0.0001,"OK","Weights do not sum to 100%")', f_bold, None, wrap_c)

    # ---------- Deals ----------
    ll = wb.create_sheet("Deals")
    put(ll, "A1", f"{len(deals)} deals that fit the {fname} thesis, scored and ranked", f_title, align=Alignment())
    put(ll, "A2", "Scores are 1-5 (blue inputs). Weights in row 4 link to the Rubric tab; rank recalculates when scores or weights change.", f_sub, align=Alignment())
    cols = ["Company", "HQ", "Sector", "Last round", "Last round ($M)", "Date",
            "Domain insight", "Prior exits", "Technical depth", "Team", "Market", "Traction", "Fund fit", "Timing",
            "Founder score (/5)", "Weighted score (/5)", "Rank", "Tier", "Notes"]
    header(ll, 5, cols, [18, 16, 26, 14, 10, 10, 9, 9, 9, 8, 8, 8, 8, 8, 10, 10, 6, 8, 52])
    put(ll, "F4", "Weights:", f_bold, align=Alignment(horizontal="right"))
    for j, col in enumerate("GHIJKLMN"):
        put(ll, f"{col}4", f"=Rubric!C{5 + j}", f_link, PCT, wrap_c)
    first, last = 6, 6 + len(deals) - 1
    for i, c in enumerate(deals):
        r = first + i
        s = c.get("scores", {})
        put(ll, f"A{r}", c["name"], f_bold)
        if c.get("website"):
            ll[f"A{r}"].hyperlink = c["website"]
        put(ll, f"B{r}", c.get("hq", ""))
        put(ll, f"C{r}", c.get("sector", ""))
        put(ll, f"D{r}", c.get("last_round", ""))
        put(ll, f"E{r}", c.get("last_round_m") if c.get("last_round_m") is not None else "Undisclosed", f_in, USD_M, wrap_c)
        put(ll, f"F{r}", c.get("last_round_date", ""), align=wrap_c)
        for col, k in zip("GHIJKLMN", KEYS):
            put(ll, f"{col}{r}", float(s.get(k, 2)), f_in, "0.0", wrap_c)
        put(ll, f"O{r}", f"=SUMPRODUCT(G{r}:J{r},$G$4:$J$4)/SUM($G$4:$J$4)", f_base, SC, wrap_c)
        put(ll, f"P{r}", f"=SUMPRODUCT(G{r}:N{r},$G$4:$N$4)", f_bold, SC, wrap_c)
        put(ll, f"Q{r}", f"=RANK(P{r},$P${first}:$P${last})", f_base, "0", wrap_c)
        put(ll, f"R{r}", f'=IF(Q{r}<=10,"Top 10","11-{len(deals)}")', f_base, None, wrap_c)
        put(ll, f"S{r}", c.get("note", ""))
        for col in range(1, 20):
            ll.cell(row=r, column=col).border = box
            if i < 10:
                ll.cell(row=r, column=col).fill = fill_band
    exc = data.get("excluded") or []
    if exc:
        er = last + 3
        put(ll, f"A{er}", "Screened out", f_bold, align=Alignment())
        for j, lab in enumerate(["Company", "HQ", "Reason"], 1):
            cc = ll.cell(row=er + 1, column=j, value=lab)
            cc.font, cc.fill, cc.border = f_head, fill_head, box
        for k, e in enumerate(exc):
            rr = er + 2 + k
            put(ll, f"A{rr}", e.get("name", ""))
            put(ll, f"B{rr}", e.get("hq", ""))
            ll.merge_cells(f"C{rr}:H{rr}")
            put(ll, f"C{rr}", e.get("reason", ""))
    ll.freeze_panes = "B6"
    ll.auto_filter.ref = f"A5:S{last}"
    LLN = f"Deals!$A${first}:$A${last}"

    def ll_lookup(col, cell):
        return f"INDEX(Deals!${col}${first}:${col}${last},MATCH({cell},{LLN},0))"

    # ---------- Market Sizing ----------
    ms = wb.create_sheet("Market Sizing")
    put(ms, "A1", "Market sizing: bottom-up TAM for the top 10", f_title, align=Alignment())
    put(ms, "A2", "TAM = base figure x filter x unit value. Blue cells are inputs; yellow marks assumptions. Top-down figures are cross-checks only.", f_sub, align=Alignment())
    header(ms, 4, ["Company", "Target customer", "Base figure", "Base figure is", "Base source", "Filter", "Filter rationale",
                   "Unit value", "Unit value rationale", "Bottom-up TAM ($M)", "Reachable share", "Reachable market ($M)",
                   "Top-down cross-check", "Cross-check source"],
           [14, 22, 16, 28, 24, 9, 30, 11, 30, 12, 10, 12, 36, 26])
    mrows = [c for c in top10 if c.get("market")]
    for i, c in enumerate(mrows):
        r = 5 + i
        m = c["market"]
        base = m.get("base", 0)
        put(ms, f"A{r}", c["name"], f_bold)
        put(ms, f"B{r}", m.get("target", ""))
        put(ms, f"C{r}", base, f_in, USD if isinstance(base, (int, float)) and base >= 1e8 else "#,##0", wrap_c)
        put(ms, f"D{r}", m.get("base_desc", ""))
        put(ms, f"E{r}", m.get("base_source", ""))
        put(ms, f"F{r}", m.get("filter", 1), f_in, PCT, wrap_c, None if m.get("filter_sourced") or m.get("filter", 1) == 1 else fill_key)
        put(ms, f"G{r}", m.get("filter_note", ""))
        put(ms, f"H{r}", m.get("unit", 0), f_in, "0.0%" if m.get("unit_is_pct") else USD, wrap_c, None if m.get("unit_sourced") else fill_key)
        put(ms, f"I{r}", m.get("unit_note", ""))
        put(ms, f"J{r}", f"=C{r}*F{r}*H{r}/1000000", f_bold, USD_M, wrap_c)
        put(ms, f"K{r}", m.get("reachable_share", 0.2), f_in, PCT, wrap_c, fill_key)
        put(ms, f"L{r}", f"=J{r}*K{r}", f_base, USD_M, wrap_c)
        put(ms, f"M{r}", m.get("crosscheck", ""))
        put(ms, f"N{r}", m.get("crosscheck_source", ""))
        boxrow(ms, r, 14)
    ms_last = 4 + max(len(mrows), 1)
    ms.freeze_panes = "B5"

    # ---------- Second Scorer ----------
    blind = [c for c in deals if c.get("blind_scores")]
    s2 = None
    if blind:
        s2 = wb.create_sheet("Second Scorer")
        put(s2, "A1", "Blind second scorer", f_title, align=Alignment())
        put(s2, "A2", "An independent reviewer scored the finalists from the same facts without seeing the first scores. Where the two differed by a point or more, the settled score below went into the Deals tab.", f_sub, align=Alignment())
        header(s2, 4, ["Company"] + KEYS + ["Blind weighted (/5)", "Final weighted (/5)", "Gap", "Blind rank", "In final top 10"],
               [22, 6, 6, 6, 6, 6, 6, 6, 6, 11, 11, 8, 8, 9])
        put(s2, "A3", "Weights:", f_bold, align=Alignment(horizontal="right"))
        for j, col in enumerate("BCDEFGHI"):
            put(s2, f"{col}3", f"=Rubric!C{5 + j}", f_link, PCT, wrap_c)
        b1, b2 = 5, 5 + len(blind) - 1
        for i, c in enumerate(blind):
            r = 5 + i
            put(s2, f"A{r}", c["name"], f_bold)
            for col, k in zip("BCDEFGHI", KEYS):
                put(s2, f"{col}{r}", float(c["blind_scores"].get(k, 2)), f_in, "0.0", wrap_c)
            put(s2, f"J{r}", f"=SUMPRODUCT(B{r}:I{r},$B$3:$I$3)", f_bold, SC, wrap_c)
            put(s2, f"K{r}", f"={ll_lookup('P', f'A{r}')}", f_link, SC, wrap_c)
            put(s2, f"L{r}", f"=K{r}-J{r}", f_base, "0.00;(0.00)", wrap_c)
            put(s2, f"M{r}", f"=RANK(J{r},$J${b1}:$J${b2})", f_base, "0", wrap_c)
            put(s2, f"N{r}", f"=IF({ll_lookup('Q', f'A{r}')}<=10,1,0)", f_base, "0", wrap_c)
            boxrow(s2, r, 14)
        r0 = b2 + 2
        put(s2, f"A{r0}", "Top-10 overlap between the two scorers", f_bold, align=Alignment())
        put(s2, f"E{r0}", f'=COUNTIFS($M${b1}:$M${b2},"<=10",$N${b1}:$N${b2},1)', f_bold, "0", wrap_c)
        put(s2, f"F{r0}", "of 10", f_base, align=Alignment())
        settled = [(c["name"], s) for c in deals for s in c.get("settled", [])]
        if settled:
            r1 = r0 + 2
            put(s2, f"A{r1}", "Settled differences", f_bold, align=Alignment())
            for j, lab in enumerate(["Company", "Factor", "First", "Blind", "Final", "Reason"], 1):
                cc = s2.cell(row=r1 + 1, column=j, value=lab)
                cc.font, cc.fill, cc.border = f_head, fill_head, box
            for k, (nm, s) in enumerate(settled):
                rr = r1 + 2 + k
                put(s2, f"A{rr}", nm)
                put(s2, f"B{rr}", s.get("factor", ""), align=wrap_c)
                put(s2, f"C{rr}", s.get("mine"), f_in, "0.0", wrap_c)
                put(s2, f"D{rr}", s.get("blind"), f_in, "0.0", wrap_c)
                put(s2, f"E{rr}", s.get("final"), f_in, "0.0", wrap_c)
                s2.merge_cells(f"F{rr}:N{rr}")
                put(s2, f"F{rr}", s.get("reason", ""))
        s2.freeze_panes = "B5"

    # ---------- Top 10 ----------
    tt = wb.create_sheet("Top 10", 0)
    put(tt, "A1", f"Top 10 deals for {fname}", f_title, align=Alignment())
    put(tt, "A2", "Ranked by founder-weighted score. Rank, scores and TAM link to other tabs (green); raised and post-money are inputs (blue).", f_sub, align=Alignment())
    tcols = ["Rank", "Company", "HQ", "Sector / model", "Founders", "Founder score (/5)", "Weighted score (/5)",
             "Blind 2nd scorer (/5)", "Funding to date", "Total raised ($M)", "Last round", "Last post-money ($M)",
             "Verified via", "Traction", "Why it fits the thesis", "Why it is interesting", "Key risks", "Next round",
             "Bottom-up TAM ($M)", "Reachable market ($M)", "Sources"]
    header(tt, 4, tcols, [6, 14, 13, 24, 44, 9, 9, 9, 34, 9, 10, 11, 16, 34, 40, 42, 32, 22, 11, 11, 32])
    for i, c in enumerate(top10):
        r = 5 + i
        put(tt, f"A{r}", f"={ll_lookup('Q', f'B{r}')}", f_link_b, "0", wrap_c)
        put(tt, f"B{r}", c["name"], f_bold)
        if c.get("website"):
            tt[f"B{r}"].hyperlink = c["website"]
        put(tt, f"C{r}", c.get("hq", ""))
        put(tt, f"D{r}", c.get("sector", ""))
        fl = "; ".join(f"{p.get('name')} ({p.get('role', '')}): {p.get('background', '')}".strip() for p in c.get("founders", []))
        put(tt, f"E{r}", fl)
        put(tt, f"F{r}", f"={ll_lookup('O', f'B{r}')}", f_link, SC, wrap_c)
        put(tt, f"G{r}", f"={ll_lookup('P', f'B{r}')}", f_link_b, SC, wrap_c)
        if s2 is not None:
            put(tt, f"H{r}", f"=IFERROR(INDEX('Second Scorer'!$J$5:$J${4 + len(blind)},MATCH(B{r},'Second Scorer'!$A$5:$A${4 + len(blind)},0)),\"-\")", f_link, SC, wrap_c)
        else:
            put(tt, f"H{r}", "-", align=wrap_c)
        put(tt, f"I{r}", c.get("funding_summary", ""))
        put(tt, f"J{r}", c.get("total_raised_m") if c.get("total_raised_m") is not None else "Undisclosed", f_in, USD_M, wrap_c)
        put(tt, f"K{r}", c.get("last_round_date", ""), align=wrap_c)
        put(tt, f"L{r}", c.get("post_money_m") if c.get("post_money_m") else "Not disclosed", f_in, USD_M, wrap_c)
        put(tt, f"M{r}", c.get("verified_via", ""))
        put(tt, f"N{r}", c.get("traction", ""))
        put(tt, f"O{r}", c.get("fit", ""))
        put(tt, f"P{r}", c.get("why", ""))
        put(tt, f"Q{r}", c.get("risks", ""))
        put(tt, f"R{r}", c.get("next_round", ""))
        if c.get("market"):
            put(tt, f"S{r}", f"=IFERROR(INDEX('Market Sizing'!$J$5:$J${ms_last},MATCH(B{r},'Market Sizing'!$A$5:$A${ms_last},0)),\"-\")", f_link, USD_M, wrap_c)
            put(tt, f"T{r}", f"=IFERROR(INDEX('Market Sizing'!$L$5:$L${ms_last},MATCH(B{r},'Market Sizing'!$A$5:$A${ms_last},0)),\"-\")", f_link, USD_M, wrap_c)
        put(tt, f"U{r}", "; ".join(c.get("sources", [])[:4]))
        boxrow(tt, r, 21)
    tt.freeze_panes = "C5"

    # ---------- Fund comps ----------
    comps = data.get("comps") or []
    facts = data.get("fund_facts") or []
    A_MED = None
    if comps or facts:
        hc = wb.create_sheet("Fund Comps")
        put(hc, "A1", f"{fname}: firm facts and recent entry prices", f_title, align=Alignment())
        put(hc, "A2", "Recent entry prices anchor the valuation on the Pick tab.", f_sub, align=Alignment())
        hc.column_dimensions["A"].width = 32
        hc.column_dimensions["B"].width = 16
        r = 4
        for f in facts:
            put(hc, f"A{r}", f.get("label", ""), f_bold)
            v = f.get("value")
            if isinstance(v, (int, float)):
                put(hc, f"B{r}", v, f_in, "#,##0.##", right)
            else:
                hc.merge_cells(f"B{r}:G{r}")
                put(hc, f"B{r}", v, f_in)
            r += 1
        if comps:
            cr = r + 1
            put(hc, f"A{cr}", "Recent entries (post-money)", f_bold, align=Alignment())
            for j, lab in enumerate(["Company", "HQ", "Round", "Date", "Round size ($M)", "Post-money ($M)", "Note"], 1):
                cc = hc.cell(row=cr + 1, column=j, value=lab)
                cc.font, cc.fill, cc.border = f_head, fill_head, box
            for col, w in zip("CDEFG", [14, 11, 13, 13, 34]):
                hc.column_dimensions[col].width = w
            for k, e in enumerate(comps):
                rr = cr + 2 + k
                put(hc, f"A{rr}", e.get("company", ""), f_bold)
                put(hc, f"B{rr}", e.get("hq", ""))
                put(hc, f"C{rr}", e.get("round", ""))
                put(hc, f"D{rr}", e.get("date", ""))
                put(hc, f"E{rr}", e.get("size_m"), f_in, USD_M, wrap_c)
                put(hc, f"F{rr}", e.get("post_m"), f_in, USD_M, wrap_c)
                put(hc, f"G{rr}", e.get("note", ""))
                boxrow(hc, rr, 7)
            ce = cr + 2 + len(comps)
            for lab, key in (("Median Series A post-money ($M)", "series a"), ("Median seed post-money ($M)", "seed")):
                cells = ",".join(f"F{cr + 2 + k}" for k, e in enumerate(comps)
                                 if str(e.get("round", "")).strip().lower() == key and e.get("post_m"))
                if cells:
                    put(hc, f"A{ce}", lab, f_bold)
                    put(hc, f"F{ce}", f"=MEDIAN({cells})", f_bold, USD_M, wrap_c)
                    if key == "series a":
                        A_MED = f"'Fund Comps'!F{ce}"
                    ce += 1

    # ---------- Pick ----------
    if pick:
        pk = wb.create_sheet("Pick", 1)
        for col, w in (("A", 3), ("B", 50), ("C", 16), ("D", 72)):
            pk.column_dimensions[col].width = w
        check = float((fund.get("check_size_usd") or 1000000))
        put(pk, "B1", f'="{pick.get("name", "")}: maximum post-money for a ${check / 1e6:,.1f}M check is $"&TEXT(C24,"0.0")&"M"', f_title, align=Alignment())
        put(pk, "B2", "Venture method: work back from a plausible exit to the highest entry price that still returns the target multiple, then compare with real prices. Blue = input, black = formula, green = link, yellow = assumption.", f_sub, align=Alignment())
        put(pk, "B4", "Why this deal", f_bold)
        put(pk, "D4", pick.get("why_won", ""))
        dil = pick.get("dilution") or {"Series A": 0.20, "Series B": 0.15, "Series C": 0.10, "Option pool": 0.05}
        dil_items = list(dil.items())[:4]
        while len(dil_items) < 4:
            dil_items.append(("(none)", 0.0))
        section(pk, 6, "Inputs")
        rows = [(7, "Check size ($)", check, USD, "From the request."),
                (8, "Entry route", pick.get("entry_route", ""), None, ""),
                (9, "Years from entry to exit", pick.get("years", 6), "0", "Assumption."),
                (10, "Exit ARR ($M)", pick.get("exit_arr_m", 75), '$#,##0"M"', "Assumption: a plausible slice of the bottom-up market."),
                (11, "Exit multiple (x ARR)", pick.get("exit_multiple", 8), MULT, "Assumption.")]
        for i, (lab, v) in enumerate(dil_items):
            rows.append((12 + i, f"Dilution: {lab}", float(v), PCT, "Assumption."))
        rows.append((16, "Target gross multiple on this check", pick.get("target_multiple", 10), MULT, "Assumption."))
        for r, lab, v, fmt, nt in rows:
            put(pk, f"B{r}", lab)
            put(pk, f"C{r}", v, f_in, fmt, Alignment(horizontal="right", vertical="top", wrap_text=True),
                fill_key if r in (10, 11, 16) else None)
            put(pk, f"D{r}", nt, f_sub)
        wait = pick.get("dilution_if_wait")
        if wait:
            wait_terms = "*".join(f"(1-{v})" for v in wait.values()) or "1"
        else:
            wait_terms = "(1-C13)*(1-C14)*(1-C15)"
        section(pk, 18, "Calculation")
        calc = [(19, "Exit value ($M)", "=C10*C11", USD_M, "Exit ARR x multiple"),
                (20, "Ownership kept through later rounds", "=(1-C12)*(1-C13)*(1-C14)*(1-C15)", PCT, ""),
                (21, "Value this check must reach at exit ($M)", "=C7*C16/1000000", USD_M, "Check x target multiple"),
                (22, "Ownership needed at exit", "=C21/C19", "0.00%", ""),
                (23, "Ownership needed at entry", "=C22/C20", "0.00%", ""),
                (24, "Maximum post-money if investing now ($M)", "=(C7/1000000)/C23", USD_M, "Walk-away price"),
                (25, "Maximum post-money if waiting for the next round ($M)", f"=(C7/1000000)/(C22/({wait_terms}))", USD_M, "Skips the next round's dilution but pays its price"),
                (26, "Implied IRR at the target multiple", "=C16^(1/C9)-1", PCT, "")]
        for r, lab, fml, fmt, nt in calc:
            b = r in (24, 25)
            put(pk, f"B{r}", lab, f_bold if b else f_base)
            put(pk, f"C{r}", fml, f_bold if b else f_base, fmt, right)
            put(pk, f"D{r}", nt, f_sub)
        section(pk, 28, "Market check")
        put(pk, "B29", "Company's last post-money ($M)")
        put(pk, "C29", pick.get("last_post_m"), f_in, USD_M, right)
        put(pk, "D29", pick.get("last_post_note", ""), f_sub)
        put(pk, "B30", "Step-up for a round now")
        put(pk, "C30", pick.get("step_up", 1.5), f_in, '0.00"x"', right, fill_key)
        put(pk, "D30", "Assumption.", f_sub)
        put(pk, "B31", "Likely price now ($M post)")
        put(pk, "C31", "=C29*C30", f_base, USD_M, right)
        put(pk, "B32", "Likely next-round price ($M post)")
        if A_MED and not pick.get("next_stage_price_m"):
            put(pk, "C32", f"={A_MED}", f_link, USD_M, right)
        else:
            put(pk, "C32", pick.get("next_stage_price_m"), f_in, USD_M, right)
        put(pk, "D32", pick.get("next_stage_note", ""), f_sub)
        put(pk, "B33", "Headroom now")
        put(pk, "C33", "=IFERROR(C24/C31-1,\"-\")", f_base, "0.0%;(0.0%)", right)
        put(pk, "B34", "Headroom at the next round")
        put(pk, "C34", "=IFERROR(C25/C32-1,\"-\")", f_base, "0.0%;(0.0%)", right)
        put(pk, "B35", "Decision", f_bold)
        put(pk, "C35", '=IF(C24>=C31,IF(C25>=C32,"Invest now or at the next round","Invest now: the next round will likely price above the cap"),"Pass or shrink the check at current prices")', f_bold)
        pk.merge_cells("C35:D35")
        put(pk, "B37", "Sensitivity: max post-money now ($M) by exit ARR (rows) and multiple (columns)", f_bold, align=Alignment())
        put(pk, "B38", "Exit ARR ($M)  \\  Multiple", f_head, fill=fill_head, align=wrap_c)
        base_arr = float(pick.get("exit_arr_m", 75))
        for j, mlt in enumerate([6, 8, 10, 12]):
            col = get_column_letter(3 + j)
            put(pk, f"{col}38", mlt, Font(name=F, size=10, bold=True, color="0000FF"), '0"x"', wrap_c, fill_band)
            pk.column_dimensions[col].width = max(pk.column_dimensions[col].width or 0, 16)
        for i, arr in enumerate([round(base_arr * f) for f in (0.67, 1, 1.33, 2)]):
            r = 39 + i
            put(pk, f"B{r}", arr, Font(name=F, size=10, bold=True, color="0000FF"), '$#,##0"M"', wrap_c, fill_band)
            for j in range(4):
                col = get_column_letter(3 + j)
                put(pk, f"{col}{r}", f"=$B{r}*{col}$38*$C$20/$C$16", f_base, USD_M, wrap_c)
        pk.column_dimensions["D"].width = 72

    # ---------- Backtest ----------
    bt_rows = data.get("backtest") or []
    if bt_rows:
        bt = wb.create_sheet("Backtest")
        put(bt, "A1", "Backtest: would the rubric have picked the fund's winners?", f_title, align=Alignment())
        put(bt, "A2", "Past first checks scored on facts knowable at entry, compared with outcomes. Win = Big win or Solid. The reviewer knew the outcomes, so treat gaps as directional.", f_sub, align=Alignment())
        ow = data.get("original_weights") or DEFAULT_WEIGHTS
        header(bt, 5, ["Company", "Outcome", "Win (1/0)"] + KEYS + ["Score, original weights", "Score, final weights", "Rank (final)", "Note"],
               [18, 12, 8, 6, 6, 6, 6, 6, 6, 6, 6, 11, 11, 8, 40])
        put(bt, "C3", "Final weights:", f_bold, align=Alignment(horizontal="right"))
        put(bt, "C4", "Original weights:", f_bold, align=Alignment(horizontal="right"))
        for j, (col, k) in enumerate(zip("DEFGHIJK", KEYS)):
            put(bt, f"{col}3", f"=Rubric!C{5 + j}", f_link, PCT, wrap_c)
            put(bt, f"{col}4", float(ow.get(k, DEFAULT_WEIGHTS[k])), f_in, PCT, wrap_c)
        f1, f2 = 6, 6 + len(bt_rows) - 1
        for i, row in enumerate(bt_rows):
            r = f1 + i
            put(bt, f"A{r}", row.get("name", ""), f_bold)
            put(bt, f"B{r}", row.get("outcome", ""))
            put(bt, f"C{r}", f'=IF(OR(B{r}="Big win",B{r}="Solid"),1,0)', f_base, "0", wrap_c)
            for col, k in zip("DEFGHIJK", KEYS):
                put(bt, f"{col}{r}", float(row.get("scores", {}).get(k, 2)), f_in, "0.0", wrap_c)
            put(bt, f"L{r}", f"=SUMPRODUCT(D{r}:K{r},$D$4:$K$4)", f_base, SC, wrap_c)
            put(bt, f"M{r}", f"=SUMPRODUCT(D{r}:K{r},$D$3:$K$3)", f_bold, SC, wrap_c)
            put(bt, f"N{r}", f"=RANK(M{r},$M${f1}:$M${f2})", f_base, "0", wrap_c)
            put(bt, f"O{r}", row.get("note", ""))
            boxrow(bt, r, 15)
        g = f2 + 2
        put(bt, f"A{g}", "Average, winners", f_bold)
        put(bt, f"A{g + 1}", "Average, others", f_bold)
        put(bt, f"A{g + 2}", "Gap", f_bold)
        for col in "DEFGHIJKLM":
            put(bt, f"{col}{g}", f"=IFERROR(AVERAGEIF($C${f1}:$C${f2},1,{col}{f1}:{col}{f2}),\"-\")", f_base, "0.00", wrap_c)
            put(bt, f"{col}{g + 1}", f"=IFERROR(AVERAGEIF($C${f1}:$C${f2},0,{col}{f1}:{col}{f2}),\"-\")", f_base, "0.00", wrap_c)
            put(bt, f"{col}{g + 2}", f"=IFERROR({col}{g}-{col}{g + 1},\"-\")", f_bold, "0.00;(0.00)", wrap_c)
        bs = data.get("backtest_summary")
        if bs:
            put(bt, f"A{g + 4}", "Signal by factor (AUC: 0.5 = none, 1.0 = perfect)", f_bold, align=Alignment())
            for j, (col, k) in enumerate(zip("DEFGHIJK", KEYS)):
                put(bt, f"{col}{g + 5}", bs["factor_auc"].get(k), f_in, "0.00", wrap_c)
            put(bt, f"A{g + 5}", "AUC", f_bold)
        bt.freeze_panes = "B6"

    # ---------- Form D leads ----------
    leads = data.get("edgar_leads") or []
    if leads:
        el = wb.create_sheet("Form D Leads")
        put(el, "A1", "SEC Form D leads", f_title, align=Alignment())
        put(el, "A2", "Recent Form D filings by likely startups in the fund's geography. Many of these raises had no press release.", f_sub, align=Alignment())
        header(el, 4, ["Company (as filed)", "City", "State", "Year inc.", "Industry", "Amount sold ($M)", "Offering total", "Filed", "Status"],
               [36, 16, 7, 9, 22, 11, 14, 11, 50])
        for i, x in enumerate(leads):
            r = 5 + i
            put(el, f"A{r}", x.get("name", ""))
            put(el, f"B{r}", (x.get("city") or "").title())
            put(el, f"C{r}", x.get("state"), align=wrap_c)
            put(el, f"D{r}", x.get("year_inc") or "", align=wrap_c)
            put(el, f"E{r}", x.get("industry"))
            put(el, f"F{r}", float(x.get("sold") or 0) / 1e6, f_in, "$#,##0.0", wrap_c)
            put(el, f"G{r}", x.get("offering_total"))
            put(el, f"H{r}", x.get("file_date"), align=wrap_c)
            put(el, f"I{r}", x.get("status", ""))
        el.freeze_panes = "B5"
        el.auto_filter.ref = f"A4:I{4 + len(leads)}"

    # ---------- Sources ----------
    srcs = data.get("sources") or []
    if srcs:
        so = wb.create_sheet("Sources")
        put(so, "A1", "Sources", f_title, align=Alignment())
        header(so, 3, ["Topic", "Source", "URL"], [16, 44, 100])
        for i, s in enumerate(srcs):
            r = 4 + i
            put(so, f"A{r}", s.get("topic", ""))
            put(so, f"B{r}", s.get("title", ""))
            c = put(so, f"C{r}", s.get("url", ""), f_url)
            if s.get("url"):
                c.hyperlink = s["url"]

    # ---------- Summary ----------
    sm = wb.create_sheet("Summary", 0)
    for col, w in (("A", 3), ("B", 7), ("C", 18), ("D", 24), ("E", 12), ("F", 12), ("G", 70)):
        sm.column_dimensions[col].width = w
    put(sm, "B1", f"{fname}: deal sourcing", f_title, align=Alignment())
    by = data.get("prepared_by")
    put(sm, "B2", ("Prepared by " + by + " | " if by else "") + str(data.get("date", "")), f_sub, align=Alignment())
    put(sm, "B4", "Bottom line", f_bold, align=Alignment())
    sm.merge_cells("B5:G5")
    put(sm, "B5", (data.get("brief") or {}).get("summary", ""))
    put(sm, "B7", "Weights", f_bold, align=Alignment())
    for i, (lab, f) in enumerate([("Founder (4 factors)", "=Rubric!C14"), ("Market", "=Rubric!C9"), ("Traction", "=Rubric!C10"),
                                  ("Fund fit", "=Rubric!C11"), ("Timing", "=Rubric!C12")]):
        r = 8 + i
        sm.merge_cells(f"B{r}:D{r}")
        put(sm, f"B{r}", lab)
        put(sm, f"E{r}", f, f_link, PCT, wrap_c)
    put(sm, "B14", "Top 10", f_bold, align=Alignment())
    for j, lab in enumerate(["Rank", "Company", "HQ", "Score (/5)", "TAM ($M)", "Why it is interesting"]):
        cc = sm.cell(row=15, column=2 + j, value=lab)
        cc.font, cc.fill, cc.alignment, cc.border = f_head, fill_head, wrap_c, box
    for i, c in enumerate(top10):
        r, tr = 16 + i, 5 + i
        put(sm, f"B{r}", f"='Top 10'!A{tr}", f_link, "0", wrap_c)
        put(sm, f"C{r}", f"='Top 10'!B{tr}", f_link_b)
        put(sm, f"D{r}", f"='Top 10'!C{tr}", f_link)
        put(sm, f"E{r}", f"='Top 10'!G{tr}", f_link, SC, wrap_c)
        put(sm, f"F{r}", f"='Top 10'!S{tr}", f_link, USD_M, wrap_c)
        put(sm, f"G{r}", c.get("why", ""))
        for col in "BCDEFG":
            sm[f"{col}{r}"].border = box
    if pick:
        put(sm, "B27", "Pick", f_bold, align=Alignment())
        for r, lab, f in ((28, "Company", pick.get("name", "")), (29, "Maximum post-money if investing now ($M)", "=Pick!C24"),
                          (30, "Likely next-round price ($M)", "=Pick!C32"), (31, "Decision", "=Pick!C35")):
            sm.merge_cells(f"B{r}:D{r}")
            put(sm, f"B{r}", lab)
            put(sm, f"E{r}", f, f_link_b if r == 29 else f_link, USD_M if r in (29, 30) else None, wrap_c if r != 31 else wrap)
            if r == 31:
                sm.merge_cells("E31:G31")
    put(sm, "B33", "Key", f_bold, align=Alignment())
    for i, (k, v, fnt) in enumerate([("Blue text", "Inputs: scores, sourced figures, assumptions.", f_in),
                                      ("Black text", "Formulas.", f_base), ("Green text", "Links to another tab.", f_link),
                                      ("Yellow fill", "Assumptions to challenge first.", f_base)]):
        r = 34 + i
        sm.merge_cells(f"B{r}:D{r}")
        cc = put(sm, f"B{r}", k, fnt)
        if k == "Yellow fill":
            cc.fill = fill_key
        sm.merge_cells(f"E{r}:G{r}")
        put(sm, f"E{r}", v)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        wrap_and_fit(ws)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
